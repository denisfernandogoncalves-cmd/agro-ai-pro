from io import BytesIO
import uuid

import pdfplumber
from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import AccessToken

from apps.accounts.models import AcessoUsuario
from apps.cadpro.models import CADPro, CADProPropriedade
from apps.financeiro.models import ParceiroFinanceiro
from apps.propriedades.models import Propriedade
from .models import FaturamentoInsumo, ProdutoEstoque


class IdentificacaoFaturamentoTests(APITestCase):
    def setUp(self):
        self.usuario = get_user_model().objects.create_user('identificacao-faturamento')
        self.client.force_authenticate(self.usuario)
        empresa = ParceiroFinanceiro.objects.create(nome='Empresa teste', tipo='fornecedor')
        produto = ProdutoEstoque.objects.create(nome='Produto teste', categoria='insumo', unidade='l')
        self.propriedade = Propriedade.objects.create(nome='Nome oficial', apelido='APELIDO_RESTRITO_A_TELA', municipio='Teste', area_hectares='24.2', cad_pro_identificacao='Principal')
        cad = CADPro.objects.create(codigo='00123', descricao='Cadastro')
        CADProPropriedade.objects.create(propriedade=self.propriedade, cad_pro=cad)
        self.dados = {'id': str(uuid.uuid4()), 'fornecedor': empresa.pk, 'produto': produto.pk, 'data_envio': '2026-10-19', 'embalagem': 'Balde', 'conteudo_embalagem': '20', 'dosagem_alqueire': '2', 'itens': [{'propriedade': self.propriedade.pk, 'cad_pro': '00123', 'area_alqueires': '10', 'quantidade_embalagens': '1'}]}

    def confirmar(self):
        resposta = self.client.post('/api/estoque/faturamentos/confirmar/', self.dados, format='json')
        self.assertEqual(resposta.status_code, 201, resposta.data)
        return resposta.data['id']

    def test_snapshot_pdf_excel_e_nome_oficial(self):
        previa = self.client.post('/api/estoque/faturamentos/previa/', self.dados, format='json')
        self.assertEqual(previa.data['itens'][0]['cad_pro_identificacao'], 'Principal')
        ident = self.confirmar()
        self.propriedade.cad_pro_identificacao = 'Secundário'
        self.propriedade.save()
        base = f'/api/estoque/faturamentos/{ident}/'
        self.assertEqual(self.client.get(base).data['resumo']['itens'][0]['cad_pro_identificacao'], 'Principal')
        repetido = self.client.post('/api/estoque/faturamentos/confirmar/', self.dados, format='json')
        self.assertEqual(repetido.status_code, 200)
        self.assertEqual(repetido.data['resumo']['itens'][0]['cad_pro_identificacao'], 'Principal')
        pdf = self.client.get(base + 'pdf/')
        self.assertEqual(pdf.status_code, 200)
        with pdfplumber.open(BytesIO(pdf.content)) as documento:
            texto = '\n'.join(p.extract_text() for p in documento.pages)
        self.assertIn('Principal', texto)
        self.assertIn('Nome oficial', texto)
        self.assertNotIn('APELIDO_RESTRITO_A_TELA', texto)
        self.assertNotIn('Secundário', texto)
        excel = self.client.get(base + 'excel/')
        self.assertEqual(excel.status_code, 200)
        wb = load_workbook(BytesIO(excel.content))
        folha = wb['Propriedades']
        self.assertEqual(folha['A2'].value, 'Nome oficial')
        self.assertEqual(folha['C2'].value, '00123')
        self.assertEqual(folha['C2'].data_type, 's')
        self.assertEqual(folha['D2'].value, 'Principal')
        self.assertIsInstance(folha['E2'].value, (int, float))
        self.assertNotIn('APELIDO_RESTRITO_A_TELA', str(list(folha.values)))

    def test_antigo_nao_inventa_identificacao_atual(self):
        ident = self.confirmar()
        registro = FaturamentoInsumo.objects.get(pk=ident)
        del registro.resumo['itens'][0]['cad_pro_identificacao']
        registro.save(update_fields=['resumo'])
        resposta = self.client.get(f'/api/estoque/faturamentos/{ident}/excel/')
        wb = load_workbook(BytesIO(resposta.content))
        self.assertEqual(wb['Propriedades']['D2'].value, 'Não informado')

    def test_excel_exige_permissao_de_imprimir(self):
        ident = self.confirmar()
        AcessoUsuario.objects.create(usuario=self.usuario, modulos=['faturamento-insumos'], permissoes={'faturamento-insumos': ['consultar']})
        self.client.force_authenticate(None)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(self.usuario)}')
        self.assertEqual(self.client.get(f'/api/estoque/faturamentos/{ident}/excel/').status_code, 403)
