from io import BytesIO
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from openpyxl import load_workbook
from .models import LoteImportacao, LinhaImportacao

class ConferenciaExportacaoTests(APITestCase):
    def test_exporta_todas_linhas_filtradas_sem_confirmar(self):
        user=get_user_model().objects.create_user("exportacao-import",password="x",is_staff=True)
        self.client.force_authenticate(user)
        lote=LoteImportacao.objects.create(arquivo_nome="audit.xlsx",arquivo_tamanho=1,arquivo_sha256="f"*64,status="com_erros",criado_por=user)
        for i in range(60):
            LinhaImportacao.objects.create(lote_importacao=lote,sequencia=i+1,planilha="Origem",linha_origem=i+2,tipo="producao",status="erro" if i<55 else "valida",hash_linha=str(i).zfill(64),erros=["Conferir código"] if i<55 else [],dados_originais={"codigo":"=1+1"})
        url=f"/api/importacoes/lotes/{lote.pk}/conferencia-excel/"
        r=self.client.get(url,{"status_linha":"erro"})
        self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content))
        self.assertEqual(wb["Linhas"].max_row,56)
        lote.refresh_from_db();self.assertEqual(lote.status,"com_erros");self.assertIsNone(lote.confirmado_em)
        self.assertEqual(self.client.get(url,{"status_linha":"inexistente"}).status_code,400)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(url).status_code,401)
