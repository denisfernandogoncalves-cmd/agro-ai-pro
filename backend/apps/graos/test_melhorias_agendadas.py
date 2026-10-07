from decimal import Decimal
from uuid import uuid4
from datetime import date, timedelta
from types import SimpleNamespace
from django.utils import timezone
from rest_framework.test import APITestCase
from .test_cargas_colhidas import CargaColhidaBase
from .models import TerceiroCadastro, MovimentoProducaoTerceiro, PendenciaConferencia, PosicaoSaldoGraos, MovimentacaoGraos
from .cargas_services import registrar_carga_colhida
from .services import registrar_ajuste, saldo_armazem
from .fechamentos import PeriodoFechado
from apps.core.auditoria import requisicao_atual
from apps.relatorios.selectors import selecionar_relatorio_operacional
from apps.core.painel import construir_painel


class MelhoriasAgendadasTests(CargaColhidaBase,APITestCase):
    def setUp(self):
        self.criar_contexto();self.client.force_authenticate(self.usuario)

    def entrada(self,nome='Externo'):
        r=self.client.post('/api/graos/terceiros/entradas/',dict(depositante=nome,cultura='Soja',safra='2026',armazem=self.armazem.pk,peso_bruto_kg='100',umidade_percentual='13',impureza_percentual='0',defeitos_percentual='0',data_entrada='2026-10-01'),format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4()))
        self.assertEqual(r.status_code,201,r.data);return r.data

    def fechar(self,cultura='Soja',safra='2026'):
        r=self.client.post('/api/core/fechamentos/',dict(armazem=self.armazem.pk,cultura=cultura,safra=safra,inicio='2026-10-01',fim='2026-10-06',justificativa='Conferência',chave_idempotencia=str(uuid4())),format='json')
        self.assertEqual(r.status_code,201,r.data)

    def test_vinculo_preserva_peso_saldo_snapshot_e_replay(self):
        e=self.entrada();t=TerceiroCadastro.objects.create(codigo='A',nome='Identidade confirmada')
        url=f"/api/graos/terceiros/entradas/{e['id']}/vincular/";d=dict(terceiro=t.pk,versao=e['versao'],motivo='Documento conferido');chave=str(uuid4())
        r=self.client.post(url,d,format='json',HTTP_IDEMPOTENCY_KEY=chave);self.assertEqual(r.status_code,201,r.data)
        self.assertEqual(r.data['peso_liquido_kg'],e['peso_liquido_kg']);self.assertEqual(r.data['saldo_kg'],e['saldo_kg'])
        m=MovimentoProducaoTerceiro.objects.get(tipo='edicao');self.assertEqual(m.delta_kg,0);self.assertEqual(m.snapshot_antes['depositante'],'Externo')
        self.assertEqual(m.snapshot_depois['terceiro'],t.pk)
        original=MovimentoProducaoTerceiro.objects.get(tipo='entrada')
        self.assertEqual(original.snapshot_depois['peso_bruto_kg'],'100.000')
        self.assertEqual(original.snapshot_depois['depositante'],'Externo')
        self.assertEqual(self.client.post(url,d,format='json',HTTP_IDEMPOTENCY_KEY=chave).status_code,200)
        self.assertEqual(self.client.post(url,d,format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4())).status_code,400)

    def test_busca_encontra_recebimento_fora_primeira_pagina(self):
        for n in range(27):self.entrada(f'Pessoa {n:03}')
        r=self.client.get('/api/graos/terceiros/entradas/',{'pagina':1,'busca':'Pessoa 000','sem_vinculo':'true'})
        self.assertEqual(r.data['total'],1);self.assertEqual(r.data['itens'][0]['depositante'],'Pessoa 000')

    def test_conciliacao_transferencia_nao_produz_area_destino(self):
        from apps.cadpro.models import CADPro, CADProPropriedade
        from apps.propriedades.models import Propriedade
        propriedade_destino=Propriedade.objects.create(nome='Outra propriedade',municipio='Sorriso',uf='MT',area_hectares='500')
        destino=CADPro.objects.create(codigo='OUTRO',descricao='Cadastro de destino');CADProPropriedade.objects.create(cad_pro=destino,propriedade=propriedade_destino)
        e=self.entrada();antes=selecionar_relatorio_operacional(secao='produtividade',pagina=1,por_pagina=25);painel=construir_painel(self.usuario);ocupacao=saldo_armazem(self.armazem)
        d=dict(versao=e['versao'],propriedade=propriedade_destino.pk,cad_pro=str(destino.pk),quantidade_kg='30',data_movimento='2026-10-02',observacoes='Titularidade confirmada')
        r=self.client.post(f"/api/graos/terceiros/entradas/{e['id']}/transferir/",d,format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4()));self.assertEqual(r.status_code,201,r.data)
        self.assertEqual(Decimal(r.data['saldo_kg']),70)
        posicao=PosicaoSaldoGraos.objects.get(propriedade=propriedade_destino,cad_pro=destino)
        self.assertEqual(posicao.saldo_fisico_kg,30)
        self.assertEqual(posicao.armazem_id,self.armazem.pk)
        depois=selecionar_relatorio_operacional(secao='produtividade',pagina=1,por_pagina=25)
        for campo in ('dados','totais_producao_propriedade','produtividade_por_cad_pro'):self.assertEqual(antes[campo],depois[campo])
        self.assertEqual(construir_painel(self.usuario)['resumo']['producao_kg'],painel['resumo']['producao_kg']);self.assertEqual(saldo_armazem(self.armazem),ocupacao)
        c=self.client.get('/api/graos/conciliacao-estoque/',dict(armazem=self.armazem.pk,cultura='Soja',safra='2026'))
        self.assertEqual(c.status_code,200,c.data);self.assertEqual(Decimal(c.data['estoque_total_kg']),100)
        self.assertEqual(Decimal(c.data['diferenca_proprio_kg']),0);self.assertEqual(Decimal(c.data['diferenca_terceiros_kg']),0)
        painel_api=self.client.get('/api/graos/saldos/painel/',{'propriedade':propriedade_destino.pk}).data
        self.assertEqual(painel_api['recebimentos_terceiros'],[{'posicao':posicao.pk,'quantidade_kg':'30.000'}])
        self.assertIn({'operacao':'ajuste','origem_externa':True,'fisico_kg':'30.000','comprometido_kg':'0.000'},painel_api['composicao'])
        self.assertEqual(sum(Decimal(c['fisico_kg']) for c in painel_api['composicao']),Decimal(painel_api['resumo']['saldo_fisico_kg']))
        movimento=MovimentoProducaoTerceiro.objects.get(tipo='transferencia')
        url=f'/api/graos/terceiros/transferencias/{movimento.pk}/'
        pdf=self.client.get(url+'pdf/');self.assertEqual(pdf.status_code,200);self.assertTrue(pdf.content.startswith(b'%PDF'))
        excel=self.client.get(url+'excel/');self.assertEqual(excel.status_code,200)
        from io import BytesIO
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(excel.content));texto=' '.join(str(c.value) for sheet in wb for row in sheet for c in row if c.value is not None)
        self.assertIn('Outra propriedade',texto);self.assertIn('OUTRO',texto);self.assertIn('Titularidade confirmada',texto);self.assertIn('não integra produção',texto)
        from django.contrib.auth import get_user_model
        from apps.accounts.models import AcessoUsuario
        consulta=get_user_model().objects.create_user('impressao-sem-transferencia')
        AcessoUsuario.objects.create(usuario=consulta,modulos=['cargas'],permissoes={'cargas':['imprimir']})
        self.client.force_authenticate(consulta);self.assertEqual(self.client.get(url+'pdf/').status_code,403)
        self.client.force_authenticate(self.usuario)
        r=self.client.post(f'/api/graos/terceiros/movimentos/{movimento.pk}/estornar/',{'motivo':'Conferência de estorno','data_movimento':'2026-10-03'},format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4()))
        self.assertEqual(r.status_code,201,r.data)
        self.assertEqual(self.client.get('/api/graos/saldos/painel/',{'propriedade':propriedade_destino.pk}).data['recebimentos_terceiros'],[])

    def test_filtro_situacao_separa_cancelados_e_saldo_esgotado(self):
        ativo=self.entrada('Ativo');cancelado=self.entrada('Cancelado');esgotado=self.entrada('Esgotado')
        r=self.client.delete(f"/api/graos/terceiros/entradas/{cancelado['id']}/",{'versao':cancelado['versao'],'motivo':'Recebimento incorreto','data_movimento':'2026-10-02'},format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4()))
        self.assertEqual(r.status_code,200,r.data)
        r=self.client.post(f"/api/graos/terceiros/entradas/{esgotado['id']}/registrar-saida/",{'quantidade_kg':'100','data_movimento':'2026-10-02','destino':'Retirada'},format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4()))
        self.assertEqual(r.status_code,201,r.data)
        for situacao,ids in [('ativos',{ativo['id'],esgotado['id']}),('cancelados',{cancelado['id']}),('esgotados',{esgotado['id']}),('',{ativo['id'],cancelado['id'],esgotado['id']})]:
            r=self.client.get('/api/graos/terceiros/entradas/',{'pagina':1,'situacao':situacao})
            self.assertEqual({i['id'] for i in r.data['itens']},ids)
        self.assertEqual(self.client.get('/api/graos/terceiros/entradas/',{'situacao':'invalida'}).status_code,400)

    def test_relatorio_pdf_excel_criterios_filtros_permissoes_e_limite(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from unittest.mock import patch
        from django.contrib.auth import get_user_model
        from apps.accounts.models import AcessoUsuario
        filtros={'secao':'estrutura','propriedade':self.propriedade.pk}
        pdf=self.client.get('/api/relatorios/operacionais/pdf/',filtros)
        self.assertEqual(pdf.status_code,200);self.assertTrue(pdf.content.startswith(b'%PDF'))
        excel=self.client.get('/api/relatorios/operacionais/exportar/',filtros)
        self.assertEqual(excel.status_code,200)
        wb=load_workbook(BytesIO(excel.content));self.assertIn('Critérios e unidades',wb.sheetnames)
        self.assertTrue(any(str(self.propriedade.pk) in str(c.value) for row in wb['Filtros'] for c in row))
        with patch('apps.relatorios.pdf.selecionar_relatorio_operacional',return_value={'dados':{'total':1001}}):
            self.assertEqual(self.client.get('/api/relatorios/operacionais/pdf/',filtros).status_code,400)
        usuario=get_user_model().objects.create_user('relatorio-sem-imprimir')
        AcessoUsuario.objects.create(usuario=usuario,modulos=['relatorios'],permissoes={'relatorios':['consultar']})
        self.client.force_authenticate(usuario)
        self.assertEqual(self.client.get('/api/relatorios/operacionais/pdf/',filtros).status_code,403)

    def test_fechamento_ajuste_rollback_confirmacao_replay(self):
        carga=registrar_carga_colhida(usuario=self.usuario,**self.dados_carga());self.fechar(safra='2026/2027')
        m=MovimentacaoGraos.objects.first();pos=m.posicao;antes=pos.saldo_fisico_kg
        request=SimpleNamespace(user=self.usuario,headers={});token=requisicao_atual.set(request)
        d=dict(usuario=self.usuario,lote=m.lote,delta_fisico_kg='10',chave_idempotencia=str(uuid4()),data_movimento=date(2026,10,2))
        try:
            with self.assertRaises(PeriodoFechado):registrar_ajuste(**d)
            pos.refresh_from_db();self.assertEqual(pos.saldo_fisico_kg,antes)
            request.headers={'Confirmar-Periodo-Fechado':'sim'};registrar_ajuste(**d)
            request.headers={};self.assertTrue(registrar_ajuste(**d).idempotente)
            pos.refresh_from_db();self.assertEqual(pos.saldo_fisico_kg,antes+10)
        finally:requisicao_atual.reset(token)

    def test_pendencias_prioridade_prazo_paginacao_duplicidade_auditada(self):
        p=PendenciaConferencia.objects.create(referencia='dup-test',tipo='duplicidade',titulo='Conferir romaneios semelhantes',detalhes={'origem':'terceiro','semelhantes':[1,2]},responsavel=self.usuario)
        d=dict(versao=1,situacao='em_analise',motivo='Conferir documentos',assumir=True,prioridade='alta',data_limite=str(timezone.localdate()-timedelta(days=1)))
        r=self.client.patch(f'/api/core/pendencias-conferencia/{p.pk}/',d,format='json');self.assertEqual(r.status_code,200,r.data);self.assertTrue(r.data['atrasada'])
        consulta=self.client.get('/api/core/pendencias-conferencia/',{'pagina':1,'situacao':'','tipo':'duplicidade','busca':'romaneios'})
        self.assertEqual(consulta.data['total'],1);self.assertEqual(consulta.data['resumo']['atrasadas'],1)
        d.update(versao=2,situacao='resolvida',motivo='Documentos diferentes: recebimentos legítimos')
        r=self.client.patch(f'/api/core/pendencias-conferencia/{p.pk}/',d,format='json');self.assertFalse(r.data['atrasada'])
        detalhe=self.client.get(f'/api/core/pendencias-conferencia/{p.pk}/').data
        self.assertEqual(len(detalhe['historico']),2);self.assertIn('legítimos',detalhe['historico'][-1]['alteracoes']['motivo'])

    def test_api_ajuste_periodo_fechado_e_cors(self):
        registrar_carga_colhida(usuario=self.usuario,**self.dados_carga());self.fechar(safra='2026/2027')
        m=MovimentacaoGraos.objects.first();d=dict(lote=m.lote_id,delta_fisico_kg='5',data_movimento='2026-10-02',chave_idempotencia=str(uuid4()))
        r=self.client.post('/api/graos/ajustes/',d,format='json');self.assertEqual(r.status_code,409,r.data);self.assertEqual(str(r.data['codigo']),'periodo_fechado')
        r=self.client.post('/api/graos/ajustes/',d,format='json',HTTP_CONFIRMAR_PERIODO_FECHADO='sim');self.assertEqual(r.status_code,201,r.data)
        self.assertEqual(self.client.post('/api/graos/ajustes/',d,format='json').status_code,200)
        from django.conf import settings
        self.assertIn('confirmar-periodo-fechado',settings.CORS_ALLOW_HEADERS)

    def test_permissoes_vinculo_conciliacao_e_prazo_nao_presumido(self):
        from django.contrib.auth import get_user_model
        from apps.accounts.models import AcessoUsuario
        e=self.entrada();t=TerceiroCadastro.objects.create(codigo='B',nome='Cadastro')
        u=get_user_model().objects.create_user('somente-consulta')
        AcessoUsuario.objects.create(usuario=u,modulos=['cargas'],permissoes={'cargas':['consultar']})
        self.client.force_authenticate(u)
        self.assertEqual(self.client.post(f"/api/graos/terceiros/entradas/{e['id']}/vincular/",dict(terceiro=t.pk,versao=e['versao'],motivo='Documento'),format='json',HTTP_IDEMPOTENCY_KEY=str(uuid4())).status_code,403)
        self.assertEqual(self.client.get('/api/graos/conciliacao-estoque/',dict(armazem=self.armazem.pk,cultura='Soja',safra='2026')).status_code,403)
        self.client.force_authenticate(self.usuario)
        p=PendenciaConferencia.objects.create(referencia='sem-prazo',tipo='peso',titulo='Peso',detalhes={},responsavel=self.usuario)
        self.assertFalse(self.client.get(f'/api/core/pendencias-conferencia/{p.pk}/').data['atrasada'])


from apps.vendas.test_saida_completa import SaidaCompletaMixin
class VendaFechadaTests(SaidaCompletaMixin,APITestCase):
    def setUp(self):self.contexto_saida();self.client.force_authenticate(self.usuario)

    def test_venda_fechada_nao_movimenta_sem_confirmar(self):
        from apps.graos.models import FechamentoPeriodo
        from apps.graos.fechamentos import movimentos_periodo,assinatura_periodo
        from apps.vendas.models import VendaGraos
        dados=self.dados.copy()
        dia=dados.get('data_saida',dados.get('data_movimento',str(timezone.localdate())))
        if isinstance(dia,str):dia=date.fromisoformat(dia)
        proprio,terceiros=movimentos_periodo(self.armazem,self.posicao.cultura,self.posicao.safra,dia)
        FechamentoPeriodo.objects.create(armazem=self.armazem,cultura=self.posicao.cultura,safra=self.posicao.safra,inicio=dia,fim=dia,justificativa='Conferência',chave_idempotencia=uuid4(),criado_por=self.usuario,saldo_proprio_kg=self.posicao.saldo_fisico_kg,saldo_terceiros_kg=0,assinatura=assinatura_periodo(proprio,terceiros))
        antes=self.posicao.saldo_fisico_kg;chave=str(uuid4())
        r=self.client.post(self.url,dados,format='json',HTTP_IDEMPOTENCY_KEY=chave);self.assertEqual(r.status_code,409,r.data)
        self.assertFalse(VendaGraos.objects.exists());self.posicao.refresh_from_db();self.assertEqual(self.posicao.saldo_fisico_kg,antes)
        r=self.client.post(self.url,dados,format='json',HTTP_IDEMPOTENCY_KEY=chave,HTTP_CONFIRMAR_PERIODO_FECHADO='sim');self.assertEqual(r.status_code,201,r.data)
