from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from rest_framework.test import APITestCase

from .cargas_services import registrar_carga_colhida, calcular_peso_liquido
from .models import EntradaProducaoTerceiro
from .test_cargas_colhidas import CargaColhidaBase


class PesagemTests(CargaColhidaBase, APITestCase):
    def setUp(self):
        self.criar_contexto()
        self.client.force_authenticate(self.usuario)
        self.dados = {"depositante": "Produtor externo", "cultura": "Soja", "safra": "2026", "armazem": self.armazem.pk, "peso_total_kg": "1500", "tara_kg": "500", "peso_bruto_kg": "9999", "umidade_percentual": "13", "impureza_percentual": "1", "defeitos_percentual": "1", "data_entrada": "2026-10-07"}

    def entrada(self, **mudancas):
        return self.client.post("/api/graos/terceiros/entradas/", {**self.dados, **mudancas}, format="json", HTTP_IDEMPOTENCY_KEY=str(uuid4()))

    def test_calculo_e_registro_usam_total_menos_tara(self):
        esperado = calcular_peso_liquido(cultura="Soja", peso_bruto_kg="1000", umidade_percentual="13", impureza_percentual="1", defeitos_percentual="1")
        previa = self.client.post("/api/graos/terceiros/previa/", self.dados, format="json")
        self.assertEqual(previa.status_code, 200, previa.data)
        self.assertEqual(Decimal(previa.data["peso_liquido_kg"]), esperado[2])
        salvo = self.entrada()
        self.assertEqual(salvo.status_code, 201, salvo.data)
        self.assertEqual(Decimal(salvo.data["peso_bruto_kg"]), Decimal(1000))
        self.assertEqual(Decimal(salvo.data["saldo_kg"]), esperado[2])
        dados = self.dados_carga()
        carga = registrar_carga_colhida(usuario=self.usuario, **{**dados, "peso_total_kg": "1500", "tara_kg": "500", "peso_bruto_kg": "9999"})
        self.assertEqual(carga.peso_bruto_kg, Decimal(1000))
        self.assertEqual(carga.peso_total_kg, Decimal(1500))

    def test_pesagens_invalidas_nao_criam_entrada(self):
        for mudancas in ({"tara_kg": None}, {"peso_total_kg": None}, {"tara_kg": "-1"}, {"tara_kg": "1500"}, {"tara_kg": "1600"}):
            self.assertEqual(self.entrada(**mudancas).status_code, 400)
        self.assertEqual(EntradaProducaoTerceiro.objects.count(), 0)

    def test_compatibilidade_com_registro_sem_tara(self):
        salvo = self.entrada(peso_total_kg=None, tara_kg=None, peso_bruto_kg="1000")
        self.assertEqual(salvo.status_code, 201, salvo.data)
        self.assertIsNone(salvo.data["tara_kg"])
        self.assertEqual(Decimal(salvo.data["peso_bruto_kg"]), Decimal(1000))

    def test_edicao_recalcula_saldo(self):
        salvo = self.entrada()
        resposta = self.client.patch(f'/api/graos/terceiros/entradas/{salvo.data["id"]}/', {**self.dados, "peso_total_kg": "2500", "versao": salvo.data["versao"], "motivo": "Correção da pesagem"}, format="json", HTTP_IDEMPOTENCY_KEY=str(uuid4()))
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(Decimal(resposta.data["peso_bruto_kg"]), Decimal(2000))
        self.assertEqual(Decimal(resposta.data["peso_liquido_kg"]), Decimal(1960))

    def test_downloads_e_planilha_sem_formulas(self):
        salvo = self.entrada(depositante="=1+1")
        pk = salvo.data["id"]
        pdf = self.client.get(f"/api/graos/terceiros/entradas/{pk}/pdf/")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        xlsx = self.client.get(f"/api/graos/terceiros/entradas/{pk}/excel/")
        self.assertEqual(xlsx.status_code, 200)
        folha = load_workbook(BytesIO(xlsx.content)).active
        valores = [c.value for linha in folha for c in linha]
        self.assertEqual(valores.count("Peso bruto do produto"), 2)
        self.assertEqual(valores.count("Responsável pelo registro"), 2)
        self.assertEqual(valores.count(self.usuario.username), 2)
        self.assertEqual(valores.count("Registrado em"), 2)
        self.assertEqual(valores.count("=1+1"), 2)
        self.assertFalse(any(c.data_type == "f" for linha in folha for c in linha))
        self.assertEqual(folha.page_setup.fitToHeight, 1)
        self.assertEqual(self.client.post(f"/api/graos/terceiros/entradas/{pk}/pdf/").status_code, 405)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(f"/api/graos/terceiros/entradas/{pk}/pdf/").status_code, (401, 403))

    def test_download_exige_permissao_de_impressao(self):
        from apps.accounts.models import AcessoUsuario
        salvo = self.entrada()
        AcessoUsuario.objects.create(usuario=self.usuario, modulos=["cargas"], permissoes={"cargas": ["consultar"]})
        self.assertEqual(self.client.get(f'/api/graos/terceiros/entradas/{salvo.data["id"]}/pdf/').status_code, 403)

    def test_observacoes_extensas_nao_sao_cortadas(self):
        salvo = self.entrada(observacoes="Observação muito extensa " * 140)
        pk = salvo.data["id"]
        pdf=self.client.get(f"/api/graos/terceiros/entradas/{pk}/pdf/")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        excel = self.client.get(f"/api/graos/terceiros/entradas/{pk}/excel/")
        self.assertEqual(excel.status_code, 200)
        folha = load_workbook(BytesIO(excel.content)).active
        self.assertTrue(any(c.value == ("Observação muito extensa " * 140).strip() for linha in folha for c in linha))
