"""Downloads do recebimento de terceiros, com duas vias para conferência."""
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle, SimpleDocTemplate, PageBreak, Spacer

from apps.vendas.romaneios import _numero


def campos_romaneio(entrada):
    def peso(valor):
        return "Não informado" if valor is None else f"{_numero(valor)} kg"

    campos = [
        ("Romaneio", f"#{entrada.pk}"),
        ("Registrado em", timezone.localtime(entrada.criado_em).strftime("%d/%m/%Y %H:%M")),
        ("Responsável pelo registro", entrada.criado_por.username),
        ("Terceiro", entrada.depositante),
        ("Data", entrada.data_entrada.strftime("%d/%m/%Y")),
        ("Produto / safra", f"{entrada.cultura} / {entrada.safra}"),
        ("Armazenagem", entrada.armazem.nome),
        ("Peso total", peso(entrada.peso_total_kg)),
        ("Tara", peso(entrada.tara_kg)),
        ("Peso bruto do produto", peso(entrada.peso_bruto_kg)),
        ("Desconto (%)", _numero(entrada.desconto_total_percentual)),
        ("Desconto (kg)", peso(entrada.desconto_total_kg)),
        ("Peso líquido", peso(entrada.peso_liquido_kg)),
        ("Sacas de 60 kg", _numero(entrada.peso_liquido_kg / Decimal(60))),
        ("Umidade (%)", _numero(entrada.umidade_percentual)),
        ("Impureza (%)", _numero(entrada.impureza_percentual)),
        ("Avariados (%)", _numero(entrada.defeitos_percentual)),
    ]
    if entrada.cultura.casefold() == "trigo":
        campos.append(("PH", _numero(entrada.ph)))
    campos.extend([
        ("Motorista / placa", f"{entrada.motorista or '—'} / {entrada.placa or 'Sem placa'}"),
        ("Documento", entrada.documento or "—"),
        ("Saldo atual", peso(entrada.saldo_kg)),
        ("Situação", "Entrada estornada" if entrada.movimentos.filter(tipo="entrada", estorno__isnull=False).exists() else "Registrada"),
        ("Observações", entrada.observacoes or "—"),
    ])
    return campos


def gerar_pdf(entrada, *, campos=None, titulo=None):
    buffer = BytesIO()
    largura, altura = A4
    pagina = canvas.Canvas(buffer, pagesize=A4)
    estilo = ParagraphStyle("romaneio", fontName="Helvetica", fontSize=8, leading=9)
    linhas = [[Paragraph(escape(k), estilo), Paragraph(escape(str(v)).replace("\n", "<br/>"), estilo)] for k, v in (campos if campos is not None else campos_romaneio(entrada))]
    tabela = Table(linhas, colWidths=[145, largura - 201])
    tabela.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, "#222222"), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    _, tamanho = tabela.wrap(largura - 56, altura / 2)
    if tamanho > altura / 2 - 95:
        completo=BytesIO()
        doc=SimpleDocTemplate(completo,pagesize=A4,leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=35)
        story=[]
        for via in ('Via do cliente','Via do arquivo'):
            if story:story.append(PageBreak())
            story.extend([Paragraph(escape(titulo or f'ENTRADA DE TERCEIROS #{entrada.pk}'),estilo),Paragraph(via,estilo),Spacer(1,10)])
            longa=Table(linhas,colWidths=[145,largura-201],splitInRow=1)
            longa.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,'#222222'),('VALIGN',(0,0),(-1,-1),'TOP')]))
            story.extend([longa,Spacer(1,20),Paragraph('Assinatura motorista: __________________    Responsável: __________________',estilo)])
        def rodape(c,documento):
            c.setFont('Helvetica',7);c.drawString(28,18,f'AGRO-AI-PRO · Conferência · Página {documento.page} · Não substitui documento fiscal.')
        doc.build(story,onFirstPage=rodape,onLaterPages=rodape)
        return completo.getvalue()
    for topo, via in ((altura - 25, "Via do cliente"), (altura / 2 - 25, "Via do arquivo")):
        pagina.setFont("Helvetica-Bold", 10)
        pagina.drawString(28, topo, titulo or f"ENTRADA DE TERCEIROS #{entrada.pk}")
        pagina.setFont("Helvetica", 9)
        pagina.drawRightString(largura - 28, topo, via)
        pagina.setFont("Helvetica", 7)
        pagina.drawString(28, topo - 14, "AGRO-AI-PRO - Resumo para conferência. Não substitui documento fiscal.")
        tabela.drawOn(pagina, 28, topo - 24 - tamanho)
        pagina.drawString(28, topo - 42 - tamanho, "Assinatura motorista: __________________    Responsável: __________________")
    pagina.setDash(3, 2)
    pagina.line(28, altura / 2, largura - 28, altura / 2)
    pagina.showPage()
    pagina.save()
    return buffer.getvalue()


def gerar_resumo_excel(dados, filtros):
    buffer = BytesIO()
    livro = Workbook()
    folha = livro.active
    folha.title = "Resumo por terceiro"
    folha.append(["Resumo de terceiros - entradas líquidas e retiradas ativas"])
    folha.append(["Gerado em", timezone.localtime().strftime("%d/%m/%Y %H:%M")])
    for campo, rotulo in (("depositante", "Terceiro"), ("cultura", "Produto"), ("safra", "Safra")):
        valor=filtros.get(campo) or 'Todos'
        if campo=='depositante':
            if filtros.get('terceiro'):valor=f"Cadastro #{filtros['terceiro']}"
            elif filtros.get('entrada'):valor=f"Recebimento #{filtros['entrada']}"
        folha.append([rotulo, valor])
        folha.cell(folha.max_row, 2).data_type = "s"
    folha.append(["Transferências para CAD/PRO estão incluídas nas saídas. Cancelamentos não são somados como entradas."])
    folha.append(["Terceiro", "Produto", "Safra", "Recebimentos ativos", "Cancelados", "Entradas (kg)", "Saídas (kg)", "Transferido CAD/PRO (kg)", "Saldo (kg)", "Cadastro / recebimento"])
    for item in dados["itens"]:
        folha.append([item["depositante"], item["cultura"], item["safra"], item["recebimentos"], item["cancelados"], *[Decimal(item[c]) for c in ("entradas_kg", "saidas_kg", "transferencias_kg", "saldo_kg")]])
        for coluna in range(1, 4):
            folha.cell(folha.max_row, coluna).data_type = "s"
        identificacao=(item.get('terceiro_codigo') or f"Cadastro #{item.get('terceiro')}") if item.get('terceiro') else f"Recebimento #{item.get('entrada_legada')} sem vínculo"
        folha.cell(folha.max_row,10,identificacao).data_type='s'
    ultima = folha.max_row
    folha.append(["TOTAL", "", "", "", "", *[Decimal(dados["totais"][c]) for c in ("entradas_kg", "saidas_kg", "transferencias_kg", "saldo_kg")]])
    folha.freeze_panes = "D8"
    folha.auto_filter.ref = f"A7:J{ultima}"
    for linha in folha.iter_rows(min_row=7):
        for celula in linha:
            celula.alignment = Alignment(wrap_text=True, vertical="top")
            if celula.row in (7, folha.max_row):
                celula.font = Font(bold=True)
            if 6 <= celula.column <= 9:
                celula.number_format = '#,##0.000'
    for coluna in "ABCDEFGHIJ":
        folha.column_dimensions[coluna].width = 25 if coluna != "A" else 34
    folha.sheet_properties.pageSetUpPr.fitToPage = True
    folha.page_setup.paperSize = folha.PAPERSIZE_A4
    folha.page_setup.orientation = "landscape"
    folha.page_setup.fitToWidth = 1
    folha.page_setup.fitToHeight = 0
    folha.print_title_rows = "1:7"
    folha.print_area = folha.dimensions
    livro.save(buffer)
    return buffer.getvalue()


def gerar_excel(entrada, *, campos=None, titulo=None):
    buffer = BytesIO()
    livro = Workbook()
    folha = livro.active
    folha.title = "Romaneio"
    borda = Border(*(Side(style="thin", color="222222") for _ in range(4)))
    for via in ("Via do cliente", "Via do arquivo"):
        folha.append([f"{titulo or f'ENTRADA DE TERCEIROS #{entrada.pk}'} - {via}"])
        folha.append(["AGRO-AI-PRO - Resumo para conferência. Não substitui documento fiscal."])
        for rotulo, valor in (campos if campos is not None else campos_romaneio(entrada)):
            folha.append([rotulo, str(valor)])
            for celula in folha[folha.max_row]:
                celula.data_type = "s"
                celula.border = borda
                celula.alignment = Alignment(wrap_text=True, vertical="top")
            folha.cell(folha.max_row, 1).font = Font(bold=True)
        folha.append(["Assinatura motorista", "Assinatura do responsável"])
        folha.append([])
    folha.column_dimensions["A"].width = 29
    folha.column_dimensions["B"].width = 72
    folha.sheet_properties.pageSetUpPr.fitToPage = True
    folha.page_setup.paperSize = folha.PAPERSIZE_A4
    folha.page_setup.fitToWidth = folha.page_setup.fitToHeight = 1
    folha.print_area = folha.dimensions
    livro.save(buffer)
    return buffer.getvalue()
