from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def _numero(valor, casas=3):
    if valor in (None, ""):
        return "—"
    try:
        numero = Decimal(valor)
    except (InvalidOperation, TypeError, ValueError):
        return str(valor)
    return f"{numero:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _dinheiro(valor):
    if valor is None:
        return "—"
    return f"R$ {_numero(valor, 2)}"


def dados_romaneio(venda, saida):
    posicao = venda.posicao
    propriedade = posicao.propriedade.nome if posicao.propriedade_id else "—"
    cadastro = f"{propriedade} / {posicao.cad_pro.codigo}"
    cultura = posicao.cultura
    qualidade = [
        ("Umidade", saida.umidade_percentual),
        ("Impureza", None),
        ("Avariado", saida.avariados_percentual),
        ("Triguilho", None),
        ("PH", saida.ph if cultura.casefold() == "trigo" else None),
    ]
    preco = venda.contrato.preco_venda if venda.contrato_id else None
    total = None
    if preco is not None:
        quantidade = Decimal(saida.quantidade_kg)
        if venda.contrato.unidade_preco == "sc":
            quantidade /= Decimal(60)
        total = _dinheiro(quantidade * Decimal(preco))
    campos = [
        ("ROMANEIO", str(saida.pk)),
        ("Data", saida.data_entrega.strftime("%d/%m/%Y")),
        ("Destinatário", saida.destino or venda.cliente_nome),
        ("Produto", cultura),
        ("Safra", posicao.safra),
        ("Propriedade / CAD/PRO", cadastro),
        ("Contrato", venda.numero_contrato or "Sem contrato"),
        ("Peso bruto", f"{_numero(saida.peso_bruto_kg)} kg"),
        ("Tara", f"{_numero(saida.tara_kg)} kg"),
        ("Peso líquido", f"{_numero(saida.quantidade_kg)} kg"),
        ("Sacas de 60 kg", _numero(Decimal(saida.quantidade_kg) / Decimal(60), 2)),
        *[(nome, f"{_numero(valor, 2)}%" if valor is not None else "—") for nome, valor in qualidade],
        ("Nota do produtor", saida.nota_produtor or "—"),
        ("Motorista", saida.motorista or "—"),
        ("Placa", saida.placa or "Sem placa"),
        ("Armazenagem", posicao.armazem.nome),
        ("Referência", saida.referencia_externa or "—"),
        ("Observações", saida.observacoes or venda.observacoes or "—"),
    ]
    if total is not None:
        campos.append(("Valor negociado", total))
    return campos


def gerar_pdf(venda, saida):
    buffer = BytesIO()
    largura, altura = A4
    pagina = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
    campos = dict(dados_romaneio(venda, saida))
    qualidade = [
        ("UMIDADE", campos.get("Umidade", "—")),
        ("IMPUREZA", campos.get("Impureza", "—")),
        ("AVARIADO", campos.get("Avariado", "—")),
        ("TRIGUILHO", campos.get("Triguilho", "—")),
        ("PH", campos.get("PH", "—")),
    ]
    separador = altura / 2

    def texto(x, y, valor, tamanho=9, negrito=False):
        pagina.setFont("Helvetica-Bold" if negrito else "Helvetica", tamanho)
        pagina.drawString(x, y, str(valor)[:95])

    def via(top, nome, arquivo=False):
        ytop = altura - top
        margem = 28
        direita = largura - margem
        texto(margem, ytop - 17, f"ROMANEIO DE SAÍDA #{saida.pk} · VENDA #{venda.pk}", 13, True)
        pagina.setFont("Helvetica-Bold", 9)
        pagina.drawRightString(direita, ytop - 17, nome.upper())
        texto(margem, ytop - 31, "AGRO-AI-PRO · Resumo para conferência. Não substitui documento fiscal.", 7)
        y = ytop - 43
        meio = largura / 2
        linhas = [
            (("ROMANEIO", campos.get("ROMANEIO")), ("DATA", campos.get("Data"))),
            (("DESTINATÁRIO", campos.get("Destinatário")), ("PRODUTO / SAFRA", f"{campos.get('Produto')} / {campos.get('Safra')}")),
            (("PROPRIEDADE / CAD/PRO", campos.get("Propriedade / CAD/PRO")), ("CONTRATO", campos.get("Contrato"))),
            (("PESO BRUTO", campos.get("Peso bruto")), ("UMIDADE", qualidade[0][1])),
            (("TARA", campos.get("Tara")), ("IMPUREZA", qualidade[1][1])),
            (("PESO LÍQUIDO", campos.get("Peso líquido")), ("AVARIADO", qualidade[2][1])),
            (("SACAS / 60 KG", campos.get("Sacas de 60 kg")), ("TRIGUILHO", qualidade[3][1])),
            (("NOTA DO PRODUTOR", campos.get("Nota do produtor")), ("PH", qualidade[4][1])),
            (("MOTORISTA", campos.get("Motorista")), ("ARMAZENAGEM", campos.get("Armazenagem"))),
            (("PLACA", campos.get("Placa")), ("REFERÊNCIA", campos.get("Referência"))),
            (("VALOR NEGOCIADO", campos.get("Valor negociado") if arquivo else "—"), ("", "")),
        ]
        altura_linha = 18
        for esquerda, direita_campo in linhas:
            for x, par in ((margem, esquerda), (meio, direita_campo)):
                pagina.setStrokeColorRGB(0.25, 0.25, 0.25)
                pagina.rect(x, y - altura_linha + 3, meio - margem, altura_linha, stroke=1, fill=0)
                if par[0]:
                    texto(x + 5, y - 9, par[0], 7, True)
                    texto(x + 105, y - 9, par[1] or "—", 8)
            y -= altura_linha
        pagina.rect(margem, y - 32, largura - 2 * margem, 32, stroke=1, fill=0)
        texto(margem + 5, y - 10, "OBSERVAÇÕES", 7, True)
        texto(margem + 5, y - 24, campos.get("Observações", "—"), 8)
        y -= 48
        pagina.rect(margem, y - 30, (largura - 2 * margem) / 2, 30, stroke=1, fill=0)
        pagina.rect(meio, y - 30, (largura - 2 * margem) / 2, 30, stroke=1, fill=0)
        texto(margem + 5, y - 23, "ASSINATURA MOTORISTA", 7, True)
        texto(meio + 5, y - 23, "ASSINATURA DO RESPONSÁVEL", 7, True)

    via(18, "Via do cliente")
    pagina.setDash(3, 2)
    pagina.line(22, separador, largura - 22, separador)
    pagina.setDash()
    pagina.setFont("Helvetica", 7)
    pagina.drawCentredString(largura / 2, separador + 4, "RECORTE ENTRE AS VIAS")
    via(separador + 14, "Via do arquivo", arquivo=True)
    pagina.showPage()
    pagina.save()
    return buffer.getvalue()


def gerar_excel(venda, saida):
    buffer = BytesIO()
    livro = Workbook()
    folha = livro.active
    folha.title = "Romaneio"
    campos = dados_romaneio(venda, saida)
    campos_cliente = [(k, v) for k, v in campos if k != "Valor negociado"]
    fino = Side(style="thin", color="000000")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    linha = 1
    for via, dados in (("VIA CLIENTE", campos_cliente), ("VIA DO ARQUIVO", campos)):
        folha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        titulo = folha.cell(linha, 1, f"ROMANEIO DE SAÍDA #{saida.pk} · VENDA #{venda.pk} — {via}")
        titulo.font = Font(bold=True, size=14, color="000000")
        linha += 1
        folha.cell(linha, 1, "AGRO-AI-PRO · Resumo para conferência. Não substitui documento fiscal.")
        folha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        linha += 1
        for rotulo, valor in dados:
            a, b = folha.cell(linha, 1, rotulo.upper()), folha.cell(linha, 2, valor)
            a.font = Font(bold=True, color="000000")
            b.font = Font(color="000000")
            a.border = b.border = borda
            a.alignment = b.alignment = Alignment(vertical="center", wrap_text=True)
            linha += 1
        folha.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        linha += 2
    folha.column_dimensions["A"].width = 28
    folha.column_dimensions["B"].width = 64
    folha.sheet_view.showGridLines = False
    folha.sheet_properties.pageSetUpPr.fitToPage = True
    folha.page_setup.orientation = "portrait"
    folha.page_setup.paperSize = folha.PAPERSIZE_A4
    folha.page_setup.fitToWidth = 1
    folha.page_setup.fitToHeight = 1
    folha.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0.1, footer=0.1)
    livro.save(buffer)
    return buffer.getvalue()
