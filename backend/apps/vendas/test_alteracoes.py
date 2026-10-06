from decimal import Decimal
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

from django.db import close_old_connections, connection
from django.test import TransactionTestCase

from rest_framework.test import APITestCase
from openpyxl import load_workbook
from apps.accounts.access import acao_da_requisicao

from apps.graos.models import MovimentacaoGraos
from apps.graos.services import creditar_producao, estornar_movimentacao, SaldoGraosError
from .alteracoes_services import editar_venda
from .services import confirmar_venda, registrar_entrega_venda, VendaGraosConflitoError
from .models import AlteracaoVendaGraos, ContratoComercial
from .tests import ContextoVendaMixin


class AlteracoesVendaTests(ContextoVendaMixin, APITestCase):
    def setUp(self):
        self.criar_contexto()
        self.client.force_authenticate(self.usuario)
        self.venda = self.rascunho()
        self.url = f"/api/comercial/vendas/{self.venda.pk}/"
        self.sequencia = 0

    def requisitar(self, metodo, caminho="", **dados):
        self.venda.refresh_from_db()
        self.sequencia += 1
        if metodo in ("patch", "delete"):
            dados = {"versao": self.venda.versao, "motivo": "Correção de teste", **dados}
        return getattr(self.client, metodo)(self.url + caminho, dados, format="json",
            HTTP_IDEMPOTENCY_KEY=f"teste-{self.sequencia}")

    def saldo(self, fisico, comprometido):
        self.posicao.refresh_from_db()
        self.assertEqual(self.posicao.saldo_fisico_kg, Decimal(fisico))
        self.assertEqual(self.posicao.saldo_comprometido_kg, Decimal(comprometido))

    def confirmar_entregar(self, quantidade="200"):
        self.assertEqual(self.requisitar("post", "confirmar/").status_code, 200)
        resposta = self.requisitar("post", "entregar/", quantidade_kg=quantidade)
        self.assertEqual(resposta.status_code, 201, resposta.data)
        return resposta.data["entregas"][0]["id"]

    def test_contrato_cadastrado_reutilizavel_e_desativacao_preserva_historico(self):
        url = "/api/comercial/contratos/"
        resposta = self.client.post(url, {"numero": "123", "empresa": "Empresa", "quantidade_kg": "1000", "produto": "Soja", "preco_venda": "125.50", "unidade_preco": "sc"}, format="json")
        self.assertEqual(resposta.status_code, 201, resposta.data)
        self.assertEqual(resposta.data["preco_venda"], "125.50")
        self.assertEqual(resposta.data["unidade_preco"], "sc")
        contrato = resposta.data["id"]
        for chave in ("primeira", "segunda"):
            resposta = self.client.post("/api/comercial/vendas/", {
                "contrato": contrato, "numero_contrato": "FALSO", "cliente_nome": "FALSO",
                "posicao": self.posicao.pk, "quantidade_kg": "100",
            }, format="json", HTTP_IDEMPOTENCY_KEY=chave)
            self.assertEqual(resposta.status_code, 201, resposta.data)
            self.assertEqual(resposta.data["numero_contrato"], "123")
            self.assertEqual(resposta.data["cliente_nome"], "Empresa")
            self.assertEqual(resposta.data["contrato_preco_venda"], "125.50")
            self.assertEqual(resposta.data["contrato_unidade_preco"], "sc")
        self.assertEqual(self.client.delete(f"{url}{contrato}/").status_code, 204)
        self.assertFalse(ContratoComercial.objects.get(pk=contrato).ativo)
        resposta = self.client.post("/api/comercial/vendas/", {
            "contrato": contrato, "posicao": self.posicao.pk, "quantidade_kg": "100",
        }, format="json", HTTP_IDEMPOTENCY_KEY="inativa")
        self.assertEqual(resposta.status_code, 400)

    def test_editar_rascunho_e_cancelada_sem_movimentar_saldo(self):
        antes = MovimentacaoGraos.objects.count()
        self.assertEqual(self.requisitar("patch", quantidade_kg="500").status_code, 200)
        self.assertEqual(self.requisitar("post", "cancelar/").status_code, 200)
        resposta = self.requisitar("patch", quantidade_kg="400")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(resposta.data["quantidade_cancelada_kg"], "400.000")
        self.assertEqual(MovimentacaoGraos.objects.count(), antes)
        self.saldo("1000", "0")


    def test_editar_confirmada_permite_reserva_acima_do_fisico(self):
        self.requisitar("post", "confirmar/")
        resposta = self.requisitar("patch", quantidade_kg="800")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("1000", "800")
        antes = MovimentacaoGraos.objects.count()
        resposta = self.requisitar("patch", quantidade_kg="1001")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertGreater(MovimentacaoGraos.objects.count(), antes)
        self.saldo("1000", "1001")

    def test_excluir_venda_estorna_entregas_devolucoes_e_reserva(self):
        self.confirmar_entregar()
        self.requisitar("post", "devolver/", quantidade_kg="50")
        self.saldo("850", "400")
        resposta = self.requisitar("delete")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertTrue(resposta.data["excluida_em"])
        self.assertTrue(resposta.data["entregas"][0]["cancelado_em"])
        self.assertTrue(resposta.data["devolucoes"][0]["cancelado_em"])
        self.saldo("1000", "0")
        self.assertEqual(self.client.get("/api/comercial/vendas/").data, [])
        self.assertEqual(len(self.client.get("/api/comercial/vendas/?mostrar_excluidas=true").data), 1)

    def test_editar_entrega_com_devolucao_preserva_totais_e_dados(self):
        entrega = self.confirmar_entregar()
        self.requisitar("post", "devolver/", quantidade_kg="50")
        resposta = self.requisitar("patch", f"entregas/{entrega}/", quantidade_kg="300", placa="ABC-1234", destino="Comprador", nota_empresa="NF123")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("750", "300")
        ativas = [e for e in resposta.data["entregas"] if not e["cancelado_em"]]
        self.assertEqual(len(ativas), 1)
        self.assertEqual(ativas[0]["placa"], "ABC1234")
        self.assertEqual(ativas[0]["nota_empresa"], "NF123")
        self.assertEqual(resposta.data["quantidade_entregue_kg"], "300.000")

    def test_excluir_entrega_com_devolucao_dependente_bloqueia(self):
        entrega = self.confirmar_entregar()
        self.requisitar("post", "devolver/", quantidade_kg="50")
        antes = MovimentacaoGraos.objects.count()
        resposta = self.requisitar("delete", f"entregas/{entrega}/")
        self.assertEqual(resposta.status_code, 409, resposta.data)
        self.assertEqual(MovimentacaoGraos.objects.count(), antes)
        self.saldo("850", "400")

    def test_editar_excluir_devolucao_e_entrega(self):
        entrega = self.confirmar_entregar()
        resposta = self.requisitar("post", "devolver/", quantidade_kg="50")
        devolucao = resposta.data["devolucoes"][0]["id"]
        resposta = self.requisitar("patch", f"devolucoes/{devolucao}/", quantidade_kg="80")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("880", "400")
        nova = next(d for d in resposta.data["devolucoes"] if not d["cancelado_em"])
        resposta = self.requisitar("delete", f"devolucoes/{nova['id']}/")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("800", "400")
        resposta = self.requisitar("delete", f"entregas/{entrega}/")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("1000", "600")

    def test_venda_entregue_pode_ser_editada_e_excluida(self):
        self.confirmar_entregar("600")
        resposta = self.requisitar("patch", observacoes="Nota corrigida")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(resposta.data["status"], "entregue")
        resposta = self.requisitar("delete")
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.saldo("1000", "0")

    def test_versao_antiga_bloqueia_e_repeticao_nao_duplica(self):
        dados = {"versao": self.venda.versao, "motivo": "Correção", "quantidade_kg": "700"}
        for _ in range(2):
            resposta = self.client.patch(self.url, dados, format="json", HTTP_IDEMPOTENCY_KEY="mesma-chave")
            self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(AlteracaoVendaGraos.objects.count(), 1)
        resposta = self.client.patch(self.url, dados, format="json", HTTP_IDEMPOTENCY_KEY="outra-chave")
        self.assertEqual(resposta.status_code, 409)
        dados["quantidade_kg"] = "800"
        resposta = self.client.patch(self.url, dados, format="json", HTTP_IDEMPOTENCY_KEY="mesma-chave")
        self.assertEqual(resposta.status_code, 409)

    def test_exclusao_repetida_e_sem_autenticacao(self):
        dados = {"versao": self.venda.versao, "motivo": "Teste encerrado"}
        for _ in range(2):
            resposta = self.client.delete(self.url, dados, format="json", HTTP_IDEMPOTENCY_KEY="excluir")
            self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(AlteracaoVendaGraos.objects.count(), 1)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.delete(self.url, dados, format="json").status_code, 401)
        self.assertEqual(self.client.get("/api/comercial/contratos/").status_code, 401)

    def test_catalogo_exige_campos_quantidade_positiva_e_permite_editar(self):
        url = "/api/comercial/contratos/"
        self.assertEqual(self.client.post(url, {"empresa": "E", "numero": "1"}).status_code, 400)
        dados = {"empresa": "E", "numero": "1", "produto": "Milho", "quantidade_kg": "0"}
        self.assertEqual(self.client.post(url, dados).status_code, 400)
        dados["quantidade_kg"] = "35000.500"
        resposta = self.client.post(url, dados)
        self.assertEqual(resposta.status_code, 201, resposta.data)
        self.assertEqual(resposta.data["quantidade_kg"], "35000.500")
        resposta = self.client.patch(f"{url}{resposta.data['id']}/", {"produto": "Soja", "quantidade_kg": "20000.250"})
        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(resposta.data["produto"], "Soja")
        self.assertEqual(resposta.data["quantidade_kg"], "20000.250")

    def test_capacidade_insuficiente_reverte_exclusao_inteira(self):
        self.confirmar_entregar("600")
        self.armazem.capacidade_kg = Decimal("1000")
        self.armazem.save()
        creditar_producao(usuario=self.usuario, lote=self.lote, quantidade_kg="600", chave_idempotencia="reposicao")
        antes = MovimentacaoGraos.objects.count()
        resposta = self.requisitar("delete")
        self.assertEqual(resposta.status_code, 409, resposta.data)
        self.assertEqual(MovimentacaoGraos.objects.count(), antes)
        self.venda.refresh_from_db()
        self.assertIsNone(self.venda.excluida_em)
        self.saldo("1000", "0")

    def test_estorno_generico_nao_desvincula_reserva_ou_entrega_da_venda(self):
        self.confirmar_entregar()
        movimentos = MovimentacaoGraos.objects.filter(reserva__isnull=False)
        self.assertEqual(movimentos.count(), 2)
        for movimento in movimentos:
            with self.subTest(movimento=movimento.pk), self.assertRaises(SaldoGraosError):
                estornar_movimentacao(usuario=self.usuario, movimentacao=movimento, chave_idempotencia=f"estorno-{movimento.pk}")
        self.saldo("800", "400")


class AlteracaoVendaConcorrenciaTests(ContextoVendaMixin, TransactionTestCase):
    def test_mesma_versao_so_permite_uma_correcao_de_saldo(self):
        if connection.vendor != "postgresql":
            self.skipTest("Concorrência exige PostgreSQL real.")
        self.criar_contexto()
        venda = self.rascunho()
        venda = confirmar_venda(usuario=self.usuario, venda=venda, chave_idempotencia="confirmar")
        barreira = Barrier(2)

        def alterar(quantidade):
            close_old_connections()
            try:
                barreira.wait(timeout=10)
                editar_venda(usuario=self.usuario, venda=venda, quantidade_kg=quantidade,
                    versao=venda.versao, motivo="Correção concorrente", chave_idempotencia=f"editar-{quantidade}")
                return "ok"
            except VendaGraosConflitoError:
                return "conflito"
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as executor:
            resultados = list(executor.map(alterar, ("700", "800")))
        self.assertCountEqual(resultados, ["ok", "conflito"])
        self.assertEqual(AlteracaoVendaGraos.objects.count(), 1)
        self.posicao.refresh_from_db()
        self.assertIn(self.posicao.saldo_comprometido_kg, (Decimal("700"), Decimal("800")))


class DownloadRomaneioTests(ContextoVendaMixin, APITestCase):
    def setUp(self):
        self.criar_contexto()
        self.client.force_authenticate(self.usuario)
        contrato = ContratoComercial.objects.create(
            numero="CTR-PDF", empresa="Cliente do teste", preco_venda="1.20", unidade_preco="kg"
        )
        self.venda = self.rascunho(numero=contrato.numero, quantidade="900")
        self.venda.contrato = contrato
        self.venda.save(update_fields=("contrato",))
        confirmar_venda(usuario=self.usuario, venda=self.venda, chave_idempotencia="confirma-romaneio")
        self.saida = registrar_entrega_venda(
            usuario=self.usuario,
            venda=self.venda,
            quantidade_kg="850",
            peso_bruto_kg="4080",
            tara_kg="3230",
            chave_idempotencia="saida-romaneio",
            destino="Cliente do teste",
        )
        self.url = f"/api/comercial/vendas/{self.venda.pk}/entregas/{self.saida.pk}"

    def test_baixar_pdf_a4_com_duas_vias(self):
        resposta = self.client.get(f"{self.url}/pdf/")
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta["Content-Type"], "application/pdf")
        self.assertTrue(resposta.content.startswith(b"%PDF-"))
        self.assertIn(b"/Count 1", resposta.content)
        self.assertIn(b"VIA DO CLIENTE", resposta.content)
        self.assertIn(b"VIA DO ARQUIVO", resposta.content)
        self.assertIn(b"R$ 1.020,00", resposta.content)

    def test_baixar_excel_com_cliente_sem_preco_e_arquivo_com_total(self):
        resposta = self.client.get(f"{self.url}/excel/")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta["Content-Type"])
        livro = load_workbook(BytesIO(resposta.content), data_only=True)
        valores = [celula.value for linha in livro.active.iter_rows() for celula in linha if celula.value]
        indice_arquivo = next(i for i, valor in enumerate(valores) if "VIA DO ARQUIVO" in str(valor))
        self.assertTrue(any("VIA CLIENTE" in str(valor) for valor in valores))
        self.assertFalse(any("VALOR NEGOCIADO" in str(valor) for valor in valores[:indice_arquivo]))
        self.assertIn("R$ 1.020,00", valores[indice_arquivo:])
        self.assertEqual(livro.active.page_setup.fitToHeight, 1)

    def test_preco_por_saca_e_download_usa_acao_imprimir(self):
        self.venda.contrato.unidade_preco = "sc"
        self.venda.contrato.preco_venda = "120.00"
        self.venda.contrato.save(update_fields=("unidade_preco", "preco_venda"))
        resposta = self.client.get(f"{self.url}/excel/")
        livro = load_workbook(BytesIO(resposta.content), data_only=True)
        valores = [celula.value for linha in livro.active.iter_rows() for celula in linha if celula.value]
        self.assertIn("R$ 1.700,00", valores)
        self.assertEqual(acao_da_requisicao(f"/api/comercial/vendas/{self.venda.pk}/entregas/{self.saida.pk}/pdf/", "GET"), "imprimir")
        self.assertEqual(acao_da_requisicao(f"/api/comercial/vendas/{self.venda.pk}/entregas/{self.saida.pk}/excel/", "GET"), "imprimir")

    def test_nao_baixa_saida_de_outra_venda(self):
        outra = self.rascunho(numero="OUTRA")
        resposta = self.client.get(f"/api/comercial/vendas/{outra.pk}/entregas/{self.saida.pk}/pdf/")
        self.assertEqual(resposta.status_code, 404)
