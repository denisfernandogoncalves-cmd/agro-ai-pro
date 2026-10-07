"""Exportação de consulta; nunca restaura ou modifica os dados exportados."""
import json
from contextlib import contextmanager
from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from django.apps import apps
from django.db import connection, transaction, models
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import permissions, serializers
from rest_framework.exceptions import PermissionDenied
from rest_framework.views import APIView

from apps.accounts.views import NoStoreResponseMixin

MAX_LINHAS = 100000
APPS_NEGOCIO = ("propriedades", "talhoes", "cadpro", "graos", "vendas", "estoque", "financeiro", "producao", "maquinas", "clima", "mercado", "importacoes")


@contextmanager
def consulta_consistente():
    ja_em_transacao = connection.in_atomic_block
    with transaction.atomic():
        if connection.vendor == "postgresql" and not ja_em_transacao:
            with connection.cursor() as cursor:
                cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
        yield


def valor_excel(valor):
    if isinstance(valor, datetime):
        return valor.astimezone(dt_timezone.utc).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    if isinstance(valor, UUID):
        return str(valor)
    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False, default=str)
    if isinstance(valor, Decimal):
        # Excel representa apenas 15 dígitos significativos com exatidão.
        return valor if len(valor.as_tuple().digits) <= 15 else str(valor)
    return valor


class PlanilhaExportacao:
    def __init__(self, descricao):
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self.linhas = 0
        self.longos = []
        self.adicionar("Informações", ["Item", "Valor"], [
            ["Conteúdo", descricao], ["Gerado em (UTC)", timezone.now()],
            ["Limite", "Cópia para consulta. Não restaura o banco. Arquivos e credenciais não incluídos."],
            ["Precisão", "Identificadores e números com mais de 15 dígitos são texto. Textos extensos estão em Textos longos."],
        ])

    def adicionar(self, nome, colunas, linhas):
        folha = self.workbook.create_sheet(nome[:31])
        folha.append(colunas)
        for celula in folha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="176348")
        for valores in linhas:
            self.linhas += 1
            if self.linhas > MAX_LINHAS:
                raise serializers.ValidationError("A exportação excede 100.000 linhas. Reduza o período ou os filtros; use o backup completo para volumes maiores.")
            folha.append([None] * len(colunas))
            for indice, valor in enumerate(valores, 1):
                celula = folha.cell(folha.max_row, indice)
                valor = valor_excel(valor)
                if isinstance(valor, str) and len(valor) > 32767:
                    referencia = f"{folha.title}!{celula.coordinate}"
                    self.longos.extend([referencia, parte // 30000 + 1, valor[parte:parte+30000]] for parte in range(0, len(valor), 30000))
                    valor = f"Texto completo em Textos longos: {referencia}"
                try:
                    celula.value = valor
                except Exception as erro:
                    raise serializers.ValidationError("Há texto com caracteres incompatíveis com Excel. Use o backup completo.") from erro
                if isinstance(valor, str):
                    celula.data_type = "s"  # não executar =, +, -, @ como fórmula.
                elif isinstance(valor, datetime):
                    celula.number_format = "dd/mm/yyyy hh:mm:ss"
                elif isinstance(valor, date):
                    celula.number_format = "dd/mm/yyyy"
                elif isinstance(valor, (Decimal, float)):
                    celula.number_format = "#,##0.###"
        folha.freeze_panes = "A2"
        folha.auto_filter.ref = folha.dimensions
        for coluna in folha.columns:
            letra = coluna[0].column_letter
            folha.column_dimensions[letra].width = min(45, max(16, len(str(coluna[0].value or "")) + 3))
        return folha

    def resposta(self, nome):
        if self.longos:
            self.adicionar("Textos longos", ["Referência", "Parte", "Texto"], self.longos)
        arquivo = BytesIO()
        self.workbook.save(arquivo)
        resposta = HttpResponse(arquivo.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resposta["Content-Disposition"] = f'attachment; filename="{nome}-{timezone.now():%Y%m%d-%H%M%S}.xlsx"'
        resposta["Cache-Control"] = "no-store, private"
        resposta["X-Content-Type-Options"] = "nosniff"
        return resposta


def modelos_backup():
    return [modelo for app in APPS_NEGOCIO for modelo in apps.get_app_config(app).get_models() if not modelo._meta.proxy]


class BackupExcelView(NoStoreResponseMixin, APIView):
    permission_classes = (permissions.IsAdminUser,)

    def get(self, request):
        if not request.user.is_active:
            raise PermissionDenied()
        planilha = PlanilhaExportacao("Backup em Excel dos dados de negócio, incluindo históricos e registros cancelados.")
        manifesto = []
        with consulta_consistente():
            for numero, modelo in enumerate(modelos_backup(), 1):
                campos = [campo for campo in modelo._meta.concrete_fields if not isinstance(campo, (models.BinaryField, models.FileField))]
                nome = f"{numero:02}_{modelo._meta.model_name}"[:31]
                def registros():
                    for item in modelo._base_manager.order_by(modelo._meta.pk.name).values_list(*(c.attname for c in campos)).iterator(chunk_size=500):
                        yield [str(v) if v is not None and (c.primary_key or c.is_relation) else v for c, v in zip(campos, item)]
                folha = planilha.adicionar(nome, [c.attname for c in campos], registros())
                manifesto.append([nome, modelo._meta.label, folha.max_row-1])
            from .models import AnexoLancamento
            metadados = AnexoLancamento.objects.order_by("id").values_list("id", "entidade", "registro_id", "nome", "tipo", "tamanho", "sha256", "criado_em", "excluido_em")
            planilha.adicionar("Documentos (metadados)", ["ID", "Entidade", "Registro", "Nome", "Tipo", "Bytes", "SHA256", "Criado em", "Excluído em"], metadados)
            planilha.adicionar("Índice", ["Aba", "Tabela", "Registros"], manifesto)
        return planilha.resposta("agro-backup-excel")
