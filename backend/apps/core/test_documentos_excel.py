from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework.test import APITestCase
from apps.accounts.models import AcessoUsuario
from apps.financeiro.models import LancamentoFinanceiro
from apps.propriedades.models import Propriedade
from apps.graos.test_cargas_colhidas import CargaColhidaBase
from apps.graos.cargas_services import registrar_carga_colhida
from apps.graos.models import CargaColhida, MovimentacaoGraos, PosicaoSaldoGraos
from apps.graos.tests import GraosSaldoBase
from apps.graos.services import estornar_movimentacao
from apps.vendas.tests import ContextoVendaMixin
from apps.vendas.services import confirmar_venda, registrar_entrega_venda
from .excel import PlanilhaExportacao
from .models import AnexoLancamento, RegistroAlteracao, RascunhoFormulario

PDF = b"%PDF-1.4\n1 0 obj << /Type /Catalog >> endobj\n%%EOF"


class ValoresExcelTests(SimpleTestCase):
    def test_textos_perigosos_ids_datas_decimais_e_texto_longo(self):
        livro = PlanilhaExportacao("Validação")
        texto = "produção " * 5000
        livro.adicionar("Dados", ["Fórmula", "ID", "Quantidade", "Data", "Precisão", "Texto"], [["=HYPERLINK(\"malicioso\")", "00123", Decimal("14.4"), date(2026,10,3), Decimal("1234567890123456.123"), texto]])
        resposta = livro.resposta("teste")
        wb = load_workbook(BytesIO(resposta.content))
        self.assertEqual(wb["Dados"]["A2"].data_type, "s")
        self.assertEqual(wb["Dados"]["B2"].value, "00123")
        self.assertEqual(wb["Dados"]["C2"].value, 14.4)
        self.assertEqual(wb["Dados"]["D2"].number_format, "dd/mm/yyyy")
        self.assertEqual(wb["Dados"]["E2"].value, "1234567890123456.123")
        partes = list(wb["Textos longos"].values)[1:]
        self.assertEqual("".join(linha[2] for linha in partes), texto)
        self.assertIn("no-store", resposta["Cache-Control"])

    def test_limite_falha_sem_exportar_planilha_parcial(self):
        from rest_framework.exceptions import ValidationError
        wb = PlanilhaExportacao("Limite")
        with patch("apps.core.excel.MAX_LINHAS", 5):
            with self.assertRaises(ValidationError):
                wb.adicionar("Dados", ["Valor"], [[1],[2]])


class DocumentosExcelTests(APITestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user("admin-excel-teste", is_staff=True)
        self.usuario = get_user_model().objects.create_user("consulta-excel-teste")
        self.acesso = AcessoUsuario.objects.create(usuario=self.usuario, modulos=["financeiro","relatorios"], permissoes={"financeiro":["consultar"],"relatorios":["consultar"]})
        self.p = Propriedade.objects.create(nome="=Fazenda", municipio="Cidade", uf="PR", area_hectares="24.2")
        self.lancamento = LancamentoFinanceiro.objects.create(tipo="pagar", descricao="Nota de teste", valor="100.50", data_vencimento="2026-10-03", propriedade=self.p)
        self.url = f"/api/core/anexos/financeiro/{self.lancamento.pk}/"
        self.client.force_authenticate(self.admin)

    def arquivo(self, nome="nota.pdf", conteudo=PDF):
        return SimpleUploadedFile(nome,conteudo,content_type="application/pdf")

    def test_backup_privado_admin_indice_e_sem_credenciais(self):
        RascunhoFormulario.objects.create(usuario=self.admin, contexto="vendas", dados={"privado":"não exportar"})
        r = self.client.get("/api/core/backup-excel/")
        self.assertEqual(r.status_code,200)
        wb=load_workbook(BytesIO(r.content))
        self.assertIn("Índice",wb.sheetnames)
        rotulos=[linha[1] for linha in list(wb["Índice"].values)[1:]]
        self.assertIn("graos.MovimentacaoGraos",rotulos)
        self.assertFalse(any(nome.startswith(("auth.","accounts.","core.")) for nome in rotulos))
        self.assertFalse(any("password" in str(v) or "não exportar" in str(v) for folha in wb for linha in folha.values for v in linha))
        folha=next(wb[linha[0]] for linha in list(wb["Índice"].values)[1:] if linha[1]=="propriedades.Propriedade")
        cabecalho=list(folha.values)[0];linha=list(folha.values)[1]
        self.assertIsInstance(linha[cabecalho.index("id")],str)
        self.assertEqual(folha.cell(2,cabecalho.index("nome")+1).data_type,"s")
        self.assertIn("no-store",r["Cache-Control"])
        self.client.force_authenticate(self.usuario)
        self.assertEqual(self.client.get("/api/core/backup-excel/").status_code,403)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get("/api/core/backup-excel/").status_code,401)

    def test_relatorio_exporta_todas_paginas_filtros_numeros_datas(self):
        for i in range(34):
            LancamentoFinanceiro.objects.create(tipo="pagar",descricao=f"Nota {i}",valor="10.25",data_vencimento="2026-10-03",propriedade=self.p)
        LancamentoFinanceiro.objects.create(tipo="pagar",descricao="Outro período",valor="1",data_vencimento="2027-10-03",propriedade=self.p)
        r=self.client.get("/api/relatorios/operacionais/exportar/",{"secao":"financeiro","data_inicio":"2026-10-03","data_fim":"2026-10-03","pagina":2,"por_pagina":1})
        self.assertEqual(r.status_code,200,getattr(r,"data",None))
        wb=load_workbook(BytesIO(r.content));linhas=list(wb["Resultados"].values)
        self.assertEqual(len(linhas)-1,35)
        self.assertEqual(linhas[1][linhas[0].index("id")],str(self.lancamento.pk))
        self.assertIsInstance(linhas[1][linhas[0].index("valor")],(int,float))
        self.assertIn("Filtros",wb.sheetnames)
        self.client.force_authenticate(self.usuario)
        self.assertEqual(self.client.get("/api/relatorios/operacionais/exportar/",{"secao":"financeiro"}).status_code,403)
        self.acesso.permissoes["relatorios"].append("imprimir");self.acesso.save()
        self.assertEqual(self.client.get("/api/relatorios/operacionais/exportar/",{"secao":"financeiro"}).status_code,200)
        self.assertEqual(self.client.get("/api/relatorios/operacionais/exportar/",{"data_inicio":"2027-01-01","data_fim":"2026-01-01"}).status_code,400)

    def test_anexo_upload_reenvio_download_exclusao_logica_auditada(self):
        for status in (201,200):
            r=self.client.post(self.url,{"arquivo":self.arquivo()},format="multipart")
            self.assertEqual(r.status_code,status,r.data)
        self.assertEqual(AnexoLancamento.objects.count(),1)
        anexo=AnexoLancamento.objects.get()
        self.assertEqual(RegistroAlteracao.objects.filter(entidade="documento do lançamento",acao="cadastrar").count(),1)
        r=self.client.get(self.url);self.assertNotIn("conteudo",r.data[0]);self.assertNotIn("url",r.data[0])
        url=f"/api/core/anexos/arquivo/{anexo.pk}/"
        r=self.client.get(url)
        self.assertEqual(b"".join(r.streaming_content),PDF)
        self.assertIn("attachment",r["Content-Disposition"])
        self.assertEqual(r["X-Content-Type-Options"],"nosniff")
        self.assertEqual(self.client.delete(url).status_code,204)
        anexo.refresh_from_db();self.assertIsNotNone(anexo.excluido_em)
        self.assertEqual(bytes(anexo.conteudo),PDF)
        self.assertEqual(self.client.get(url).status_code,404)
        self.assertEqual(self.client.get(self.url).data,[])
        self.assertTrue(RegistroAlteracao.objects.filter(entidade="documento do lançamento",acao="excluir").exists())
        self.assertEqual(self.client.post(self.url,{"arquivo":self.arquivo()},format="multipart").status_code,201)

    def test_anexo_permissoes_separadas_e_restricao_do_modulo(self):
        r=self.client.post(self.url,{"arquivo":self.arquivo()},format="multipart");pk=r.data["id"]
        self.client.force_authenticate(self.usuario)
        self.assertEqual(self.client.get(self.url).status_code,200)
        self.assertEqual(self.client.post(self.url,{"arquivo":self.arquivo()},format="multipart").status_code,403)
        self.assertEqual(self.client.delete(f"/api/core/anexos/arquivo/{pk}/").status_code,403)
        self.acesso.modulos=["relatorios"];self.acesso.save()
        self.assertEqual(self.client.get(self.url).status_code,403)
        self.assertEqual(self.client.get(f"/api/core/anexos/arquivo/{pk}/").status_code,403)
        self.assertIsNone(AnexoLancamento.objects.get(pk=pk).excluido_em)

    def test_anexo_invalido_grande_tipo_nome_alvo(self):
        for nome, conteudo in [("arquivo.html",b"<html/>"),("falso.pdf",b"<script/>"),("arquivo.pdf",PDF+b"a"*(5*1024*1024))]:
            r=self.client.post(self.url,{"arquivo":self.arquivo(nome,conteudo)},format="multipart")
            self.assertEqual(r.status_code,400,r.data)
        self.assertEqual(AnexoLancamento.objects.count(),0)
        self.assertEqual(self.client.get("/api/core/anexos/desconhecido/1/").status_code,400)
        self.assertEqual(self.client.get("/api/core/anexos/financeiro/invalido/").status_code,400)
        self.assertEqual(self.client.get("/api/core/anexos/financeiro/99999/").status_code,404)

    def test_anexo_uuid_canonico_sem_duplicar_por_maiusculas(self):
        from apps.estoque.models import FaturamentoInsumo, ProdutoEstoque
        from apps.financeiro.models import ParceiroFinanceiro
        fornecedor=ParceiroFinanceiro.objects.create(nome="Fornecedor teste",tipo="fornecedor")
        produto=ProdutoEstoque.objects.create(nome="Produto teste",categoria="insumo",unidade="l")
        faturamento=FaturamentoInsumo.objects.create(fornecedor=fornecedor,produto=produto,data_envio=date(2026,10,3),criado_por=self.admin,assinatura="teste",resumo={})
        for registro,status in [(str(faturamento.pk).upper(),201),(str(faturamento.pk),200)]:
            r=self.client.post(f"/api/core/anexos/faturamento/{registro}/",{"arquivo":self.arquivo()},format="multipart")
            self.assertEqual(r.status_code,status,r.data)
            self.assertEqual(r.data["registro_id"],str(faturamento.pk))
        self.assertEqual(AnexoLancamento.objects.count(),1)

    def test_possivel_duplicata_financeiro_so_consulta_e_restricoes(self):
        dados={"data":"2026-10-03","quantidade":"100.500","descricao":"Nota de teste"}
        r=self.client.post("/api/core/duplicidades/financeiro/",dados,format="json")
        self.assertEqual(r.status_code,200,r.data);self.assertEqual(r.data["total"],1)
        self.assertEqual(LancamentoFinanceiro.objects.count(),1)
        dados["quantidade"]="100.501"
        self.assertEqual(self.client.post("/api/core/duplicidades/financeiro/",dados,format="json").data["total"],0)
        self.client.force_authenticate(self.usuario)
        self.assertEqual(self.client.post("/api/core/duplicidades/venda/",dados,format="json").status_code,403)


class DuplicidadeCargaTests(CargaColhidaBase,APITestCase):
    def setUp(self):
        self.criar_contexto();self.client.force_authenticate(self.usuario)
        self.carga=registrar_carga_colhida(usuario=self.usuario,**self.dados_carga())

    def test_normaliza_placa_exclui_proprio_registro_e_nao_grava(self):
        d={"data":"2026-08-09","quantidade":"1000","placa":"ABC-1D23","propriedade":self.propriedade.pk,"cultura":"Soja","safra":"2026/2027"}
        r=self.client.post("/api/core/duplicidades/carga/",d,format="json")
        self.assertEqual(r.status_code,200,r.data);self.assertEqual(r.data["total"],1)
        d["excluir_id"]=self.carga.pk
        self.assertEqual(self.client.post("/api/core/duplicidades/carga/",d,format="json").data["total"],0)
        self.assertEqual(CargaColhida.objects.count(),1)


class HistoricoFiltradoTests(GraosSaldoBase,APITestCase):
    def setUp(self):
        self.criar_contexto();self.client.force_authenticate(self.usuario)
        credito=self.creditar("100");self.posicao=PosicaoSaldoGraos.objects.get()
        estornar_movimentacao(usuario=self.usuario,movimentacao=credito.movimentacoes[0],chave_idempotencia="estorno-historico",observacoes="Teste")
        self.creditar("25","credito-seguinte")

    def test_filtro_mantem_saldos_acumulados_e_contexto(self):
        r=self.client.get(f"/api/core/conferencia/{self.posicao.pk}/",{"operacao":"credito_producao"})
        self.assertEqual(r.status_code,200,r.data)
        self.assertEqual(r.data["total_movimentos"],2);self.assertEqual(r.data["total_historico"],3)
        self.assertEqual([Decimal(i["saldo_acumulado_kg"]) for i in r.data["itens"]],[100,25])
        self.assertEqual(r.data["cultura"],"Soja")
        self.assertEqual(Decimal(r.data["totais"]["fisico"]),25)
        self.posicao.refresh_from_db();self.assertEqual(self.posicao.saldo_fisico_kg,25)
        self.assertEqual(self.client.get(f"/api/core/conferencia/{self.posicao.pk}/",{"operacao":"desconhecida"}).status_code,400)


class DuplicidadeVendaTests(ContextoVendaMixin, APITestCase):
    def setUp(self):
        self.criar_contexto();self.client.force_authenticate(self.usuario)
        self.venda=self.rascunho(quantidade="100")
        confirmar_venda(usuario=self.usuario,venda=self.venda,chave_idempotencia="confirmar-documentos")
        registrar_entrega_venda(usuario=self.usuario,venda=self.venda,quantidade_kg="50",data_entrega=date(2026,10,3),destino="Empresa",placa="ABC1D23",nota_empresa="00123",chave_idempotencia="entrega-documentos")

    def test_nota_ou_data_placa_peso_e_empresa_sem_alterar_saldo(self):
        d={"data":"2026-10-03","quantidade":"50","placa":"ABC-1D23","destino":"Empresa","cultura":"Soja"}
        r=self.client.post("/api/core/duplicidades/venda/",d,format="json")
        self.assertEqual(r.status_code,200,r.data);self.assertEqual(r.data["total"],1)
        d.update(data="2026-10-04",quantidade="45",nota_empresa="00123")
        self.assertEqual(self.client.post("/api/core/duplicidades/venda/",d,format="json").data["total"],1)
        d["destino"]="Outra empresa"
        self.assertEqual(self.client.post("/api/core/duplicidades/venda/",d,format="json").data["total"],0)
        self.posicao.refresh_from_db();self.assertEqual(self.posicao.saldo_fisico_kg,950)
