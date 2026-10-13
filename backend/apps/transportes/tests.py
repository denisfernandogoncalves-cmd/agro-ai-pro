import io
from pathlib import Path
from unittest import skipUnless
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework.exceptions import PermissionDenied, ValidationError

from apps.accounts.access import verificar_acesso
from apps.accounts.models import AcessoUsuario
from .importacao import importar, ler_planilha
from .models import Caminhao, LancamentoTransporte, RemuneracaoMotorista
from .relatorios import relatorio


def arquivo(valor=100, sem_data=False):
    w = openpyxl.Workbook(); s = w.active; s.title = "JAN"
    for chave, v in {"B2": "Motorista A", "D3": "VOLVO", "D4": "ABC1D23", "F5": 2026, "B3": 1500, "B5": 200, "A9": date(2026, 1, 10), "B9": "Origem / destino", "C9": "MILHO", "D9": 2, "E9": valor, "I9": date(2026, 1, 11), "J9": "COMBUSTÍVEL", "M9": 1, "N9": 50}.items():
        s[chave] = v
    if sem_data: s["A9"] = None
    saida = io.BytesIO(); w.save(saida)
    return SimpleUploadedFile("caminhao.xlsx", saida.getvalue())


class TransportesTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="transporte", password="teste")
        self.client = APIClient(); self.client.force_authenticate(self.user)
        self.c = Caminhao.objects.create(placa="XYZ9A99", veiculo="Volvo", motorista="Motorista A")

    def lancamento(self, **kwargs):
        dados = dict(caminhao=self.c, motorista="Motorista A", data=date(2026, 1, 10), tipo="receita", descricao="MILHO", quantidade=Decimal(2), valor_unitario=Decimal(100))
        dados.update(kwargs)
        return LancamentoTransporte.objects.create(**dados)

    def test_relatorio_inclui_salario_comissao_sem_recalcular(self):
        self.lancamento(); self.lancamento(tipo="despesa", quantidade=Decimal(1), valor_unitario=Decimal(30))
        RemuneracaoMotorista.objects.create(caminhao=self.c, motorista="Motorista A", competencia=date(2026, 1, 1), salario=100, comissao=15)
        r = relatorio({"data_inicio": "2026-01-01", "data_fim": "2026-01-31"})
        self.assertEqual(r["resumo"]["saldo"], "55.00")
        self.assertEqual(r["resumo"]["comissao"], "15.00")
        self.assertEqual(r["resumo"]["margem"], "27.50")

    def test_percentual_desconta_apenas_areia_pedra_e_recalcula(self):
        self.lancamento(quantidade=Decimal(1), valor_unitario=Decimal(1000))
        self.lancamento(tipo="despesa", descricao="AREIA", quantidade=Decimal(1), valor_unitario=Decimal(100))
        self.lancamento(tipo="despesa", descricao="PEDRA", quantidade=Decimal(1), valor_unitario=Decimal(50))
        self.lancamento(tipo="despesa", descricao="COMBUSTÍVEL", quantidade=Decimal(1), valor_unitario=Decimal(200))
        rem = RemuneracaoMotorista.objects.create(caminhao=self.c, motorista="Motorista A", competencia=date(2026, 1, 1), salario=2500, comissao_percentual=Decimal(12))
        self.assertEqual(rem.valor_comissao, Decimal("102.00"))
        self.assertEqual(relatorio({})["resumo"]["comissao"], "102.00")
        self.lancamento(tipo="despesa", descricao="AREIA", quantidade=Decimal(1), valor_unitario=Decimal(1000))
        self.assertEqual(rem.valor_comissao, Decimal("0.00"))

    def test_api_exige_percentual_valido_e_nao_aceita_valor_manual(self):
        self.lancamento()
        dados = dict(caminhao=self.c.pk, motorista="Motorista A", competencia="2026-01-01", salario="2500", comissao="9999")
        self.assertEqual(self.client.post("/api/transportes/remuneracoes/", dados).status_code, 400)
        dados["comissao_percentual"] = "101"
        self.assertEqual(self.client.post("/api/transportes/remuneracoes/", dados).status_code, 400)
        dados["comissao_percentual"] = "12"
        resposta = self.client.post("/api/transportes/remuneracoes/", dados)
        self.assertEqual(resposta.status_code, 201)
        self.assertEqual(resposta.data["comissao"], "24.00")
        previa = self.client.get("/api/transportes/remuneracoes/calculo/", {**dados, "comissao_percentual": "8"})
        self.assertEqual(previa.status_code, 200)
        self.assertEqual(previa.data["comissao"], "16.00")

    def test_periodo_parcial_nao_rateia_remuneracao(self):
        RemuneracaoMotorista.objects.create(caminhao=self.c, motorista="Motorista A", competencia=date(2026, 1, 1), salario=100, comissao=15)
        self.assertEqual(relatorio({"data_fim": "2026-01-30"})["resumo"]["salario"], "0")
        self.assertEqual(relatorio({"data_inicio": "2026-01-31", "data_fim": "2026-01-31"})["resumo"]["salario"], "100.00")

    def test_multiplas_placas_ranking_e_filtro_motorista_historico(self):
        self.lancamento()
        outra = Caminhao.objects.create(placa="DEF2G34", veiculo="Scania", motorista="Motorista B")
        self.lancamento(caminhao=outra, motorista="Motorista B", valor_unitario=Decimal(300))
        self.c.motorista = "Outro motorista"; self.c.save()
        r = relatorio({})
        self.assertEqual(r["por_placa"][0]["nome"], "DEF2G34")
        self.assertEqual(relatorio({"motorista": "Motorista A"})["resumo"]["receita"], "200.00")
        self.assertEqual(relatorio({"placa": "XYZ-9A99"})["resumo"]["receita"], "200.00")

    def test_total_arredondado_meio_centavo(self):
        self.assertEqual(self.lancamento(quantidade=Decimal(1), valor_unitario=Decimal("1.005")).total, Decimal("1.01"))

    def test_previa_sem_gravacao_importacao_idempotente(self):
        dados = ler_planilha(arquivo())
        self.assertEqual(LancamentoTransporte.objects.count(), 0)
        self.assertEqual(len(dados["lancamentos"]), 2)
        self.assertEqual(importar(dados)["lancamentos_criados"], 2)
        self.assertEqual(importar(ler_planilha(arquivo()))["lancamentos_criados"], 0)
        rem = RemuneracaoMotorista.objects.get()
        self.assertEqual(rem.salario, Decimal("1500.00")); self.assertEqual(rem.comissao, Decimal("200.00"))

    def test_reimportacao_alterada_nao_sobrescreve(self):
        importar(ler_planilha(arquivo()))
        with self.assertRaises(ValidationError): ler_planilha(arquivo(valor=999))
        self.assertEqual(LancamentoTransporte.objects.get(tipo="receita").valor_unitario, Decimal(100))

    def test_lancamento_sem_data_rejeitado(self):
        with self.assertRaises(ValidationError): ler_planilha(arquivo(sem_data=True))

    def test_planilha_invalida(self):
        with self.assertRaises(ValidationError): ler_planilha(SimpleUploadedFile("a.xlsx", b"invalido"))
        with self.assertRaises(ValidationError): ler_planilha(SimpleUploadedFile("a.xls", b"invalido"))

    def test_validacao_formulario_e_competencia(self):
        r = self.client.post("/api/transportes/lancamentos/", dict(caminhao=self.c.pk, motorista="Motorista A", data="2026-01-01", tipo="receita", descricao="teste", quantidade="-1", valor_unitario="10"))
        self.assertEqual(r.status_code, 400)
        r = self.client.post("/api/transportes/remuneracoes/", dict(caminhao=self.c.pk, motorista="Motorista A", competencia="2026-01-15", salario="10", comissao="1"))
        self.assertEqual(r.status_code, 400)

    def test_filtro_invertido_rejeitado(self):
        r = self.client.get("/api/transportes/lancamentos/relatorio/?data_inicio=2026-02-01&data_fim=2026-01-01")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.client.get("/api/transportes/lancamentos/?data_inicio=abc").status_code, 400)

    def test_paginacao_nao_limita_totais(self):
        for _ in range(35): self.lancamento()
        self.assertEqual(len(self.client.get("/api/transportes/lancamentos/").data["results"]), 30)
        self.assertEqual(self.client.get("/api/transportes/lancamentos/relatorio/").data["resumo"]["receita"], "7000.00")

    def test_exportacao_inclui_todos_filtrados_e_neutraliza_formula(self):
        for _ in range(35): self.lancamento(descricao="=1+1")
        resposta = self.client.get("/api/transportes/lancamentos/exportar/?data_inicio=2026-01-01")
        self.assertEqual(resposta.status_code, 200)
        texto = resposta.content.decode("utf-8-sig")
        self.assertEqual(texto.count(";'=1+1;"), 35)
        self.assertIn("7000,00", texto)

    def test_autenticacao_e_permissoes_por_acao(self):
        self.assertEqual(APIClient().get("/api/transportes/caminhoes/").status_code, 401)
        AcessoUsuario.objects.create(usuario=self.user, modulos=["transportes"], permissoes={"transportes": ["consultar"]})
        verificar_acesso(self.user, "/api/transportes/caminhoes/", "GET")
        verificar_acesso(self.user, "/api/transportes/lancamentos/previa/", "POST")
        for url, method in (("/api/transportes/caminhoes/", "POST"), ("/api/transportes/lancamentos/importar/", "POST"), ("/api/transportes/lancamentos/exportar/", "GET")):
            with self.assertRaises(PermissionDenied): verificar_acesso(self.user, url, method)

    def test_caminhao_com_historico_protegido(self):
        self.lancamento()
        self.assertEqual(self.client.delete(f"/api/transportes/caminhoes/{self.c.pk}/").status_code, 400)
        self.assertTrue(Caminhao.objects.filter(pk=self.c.pk).exists())

    def test_placa_normalizada_e_unica(self):
        resposta = self.client.post("/api/transportes/caminhoes/", dict(placa="abc-1d23", veiculo="Volvo", motorista="B"))
        self.assertEqual(resposta.status_code, 201)
        self.assertEqual(resposta.data["placa"], "ABC1D23")
        self.assertEqual(self.client.post("/api/transportes/caminhoes/", dict(placa="ABC1D23", veiculo="Volvo", motorista="B")).status_code, 400)

    def test_tracos_vazios_do_modelo_nao_sao_lancamentos(self):
        original = arquivo(); w = openpyxl.load_workbook(io.BytesIO(original.read()))
        w.active["C53"] = "-"; w.active["J83"] = "-"
        saida = io.BytesIO(); w.save(saida)
        dados = ler_planilha(SimpleUploadedFile("teste.xlsx", saida.getvalue()))
        self.assertEqual(len(dados["lancamentos"]), 2)

    @skipUnless(Path("/tmp/transportes-reference.xlsx").exists(), "Planilha privada de referência não disponível neste ambiente")
    def test_reconciliacao_planilha_real_sem_alterar_original(self):
        from .importacao import MESES
        conteudo = Path("/tmp/transportes-reference.xlsx").read_bytes()
        dados = ler_planilha(SimpleUploadedFile("RHD 4D04.xlsx", conteudo))
        self.assertEqual(len(dados["lancamentos"]), 405)
        self.assertEqual(len(dados["remuneracoes"]), 9)
        w = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
        for item in dados["lancamentos"]:
            aba, celula = item["origem"].split("!"); linha = int(celula[1:])
            salvo = Decimal(str(w[aba].cell(linha, 6 if item["tipo"] == "receita" else 15).value))
            calculado = Decimal(item["quantidade"]) * Decimal(item["valor_unitario"])
            self.assertEqual(salvo.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), calculado.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), item["origem"])
        importar(dados)
        for i in dados["remuneracoes"]:
            mes = MESES[int(i["competencia"][5:7]) - 1]
            self.assertEqual(Decimal(i["comissao"]), Decimal(str(w[mes]["B5"].value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        self.assertEqual(importar(ler_planilha(SimpleUploadedFile("RHD 4D04.xlsx", conteudo)))["lancamentos_criados"], 0)
