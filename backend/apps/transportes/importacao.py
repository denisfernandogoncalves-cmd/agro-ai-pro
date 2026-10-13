import hashlib

import io

import json

import zipfile

from datetime import date, datetime

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP



import openpyxl

from django.core.exceptions import ValidationError as DjangoValidationError

from django.db import transaction

from rest_framework.exceptions import ValidationError



from .models import Caminhao, LancamentoTransporte, RemuneracaoMotorista

from .serializers import CaminhaoSerializer



MESES = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ")





def numero(valor, campo):

    try:

        resultado = Decimal(str(valor or 0))

        if not resultado.is_finite() or resultado < 0:

            raise ValueError()

        return resultado

    except (InvalidOperation, ValueError):

        raise ValidationError(f"{campo}: número inválido ou fórmula sem valor salvo. Abra e salve a planilha no Excel.")





def texto(valor):

    return str(valor or "").strip()





def ler_planilha(arquivo):

    if not arquivo or not arquivo.name.lower().endswith(".xlsx") or arquivo.size > 10 * 1024 * 1024:

        raise ValidationError("Envie uma planilha .xlsx de até 10 MB.")

    conteudo = arquivo.read()

    try:

        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:

            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:

                raise ValidationError("Planilha descompactada excede 50 MB.")

        valores = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)

        formulas = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=False, read_only=True)

    except (zipfile.BadZipFile, OSError, ValueError, KeyError):

        raise ValidationError("Não foi possível ler a planilha .xlsx.")

    lancamentos, remuneracoes, veiculos = [], [], {}

    try:

        if not any(nome in valores.sheetnames for nome in MESES):

            raise ValidationError("Planilha incompatível: esperadas abas JAN a DEZ.")

        for mes, nome in enumerate(MESES, 1):

            if nome not in valores.sheetnames:

                continue

            s = valores[nome]

            rows = list(s.iter_rows(min_row=1, max_row=min(s.max_row, 1003), max_col=15, values_only=True))

            cell = lambda linha, coluna: rows[linha - 1][coluna - 1] if linha <= len(rows) else None

            registros = [(i, row) for i, row in enumerate(rows[8:], 9) if row[0] or row[8]]

            for linha, row in enumerate(rows[8:], 9):

                if (any(v not in (None, "", "-") for v in row[1:5]) and not row[0]) or (any(v not in (None, "", "-") for v in row[9:14]) and not row[8]):

                    raise ValidationError(f"{nome}, linha {linha}: lançamento preenchido sem data.")

            salario, comissao = cell(3, 2), cell(5, 2)

            if not registros and not salario and not comissao:

                continue

            placa = texto(cell(4, 4)).upper().replace("-", "").replace(" ", "")

            motorista = texto(cell(2, 2))

            caminh = dict(placa=placa, veiculo=texto(cell(3, 4)) or "Caminhão", motorista=motorista)

            validador = CaminhaoSerializer(data=caminh)

            # A placa já existente será vinculada, sem sobrescrever seu cadastro.

            validador.fields["placa"].validators = [v for v in validador.fields["placa"].validators if v.__class__.__name__ != "UniqueValidator"]

            validador.is_valid(raise_exception=True)

            veiculos[placa] = caminh

            for linha, row in registros:

                for tipo, offset in (("receita", 0), ("despesa", 8)):

                    data = row[offset]

                    if not data:

                        continue

                    if not isinstance(data, (date, datetime)):

                        raise ValidationError(f"{nome}, linha {linha}: data inválida.")

                    data = data.date() if isinstance(data, datetime) else data

                    q = numero(row[offset + (3 if tipo == "receita" else 4)], f"{nome}, linha {linha}, quantidade")

                    v = numero(row[offset + (4 if tipo == "receita" else 5)], f"{nome}, linha {linha}, valor")

                    if row[offset + (4 if tipo == "receita" else 5)] is None:

                        raise ValidationError(f"{nome}, linha {linha}: valor unitário ausente.")

                    if q <= 0:

                        raise ValidationError(f"{nome}, linha {linha}: quantidade deve ser maior que zero.")

                    item = dict(placa=placa, motorista=motorista, data=data.isoformat(), tipo=tipo, descricao=texto(row[2] if tipo == "receita" else row[9]) or ("Frete" if tipo == "receita" else "Diversos"), local=texto(row[1] if tipo == "receita" else row[10]), nota=texto(row[11]) if tipo == "despesa" else "", unidade="", quantidade=str(q), valor_unitario=str(v), origem=f"{nome}!{'A' if tipo == 'receita' else 'I'}{linha}")

                    # Chave da posição e identidade; reimportações não duplicam linhas anteriores.

                    chave = json.dumps([placa, motorista, data.year, nome, linha, tipo], ensure_ascii=True)

                    item["chave_importacao"] = hashlib.sha256(chave.encode()).hexdigest()

                    instancia = LancamentoTransporte(**{k: v for k, v in item.items() if k != "placa"})

                    try:

                        instancia.clean_fields(exclude=["caminhao"])

                    except DjangoValidationError as exc:

                        raise ValidationError(f"{nome}, linha {linha}: {exc}") from exc

                    lancamentos.append(item)

            if formulas[nome]["B5"].data_type == "f" and comissao is None:

                raise ValidationError(f"{nome}!B5: comissão sem valor salvo. Abra e salve no Excel.")

            ano = cell(5, 6) or (registros[0][1][0] or registros[0][1][8]).year

            try:

                competencia = date(int(ano), mes, 1)

            except (ValueError, TypeError):

                raise ValidationError(f"{nome}: ano inválido.")

            if numero(cell(4, 2), nome + " percentual") > 100:

                raise ValidationError(f"{nome}!B4: percentual deve estar entre 0 e 100.")

            remuneracoes.append(dict(placa=placa, motorista=motorista, competencia=competencia.isoformat(), salario=str(numero(salario, nome + " salário").quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)), comissao=str(numero(comissao, nome + " comissão").quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)), base_por_aba=True, comissao_percentual=str(numero(cell(4, 2), nome + " percentual")) if cell(4, 2) is not None else None, observacoes=f"Importado de {nome}!B3/B4/B5; percentual manual e base da aba."))

    finally:

        valores.close()

        formulas.close()

    chaves = set(LancamentoTransporte.objects.filter(chave_importacao__in=[i["chave_importacao"] for i in lancamentos]).values_list("chave_importacao", flat=True))

    conflitos = []

    for i in lancamentos:

        if i["chave_importacao"] in chaves:

            atual = LancamentoTransporte.objects.get(chave_importacao=i["chave_importacao"])

            if any(str(getattr(atual, k)) != str(i[k]) for k in ("data", "descricao", "local", "nota")) or atual.quantidade != Decimal(i["quantidade"]) or atual.valor_unitario != Decimal(i["valor_unitario"]):

                conflitos.append(i["origem"])

    if conflitos:

        raise ValidationError("Linhas já importadas foram alteradas: " + ", ".join(conflitos[:10]) + ". Revise os lançamentos existentes antes de reimportar.")

    return dict(veiculos=list(veiculos.values()), lancamentos=lancamentos, remuneracoes=remuneracoes, duplicados=len(chaves))





@transaction.atomic

def importar(dados):

    caminh = {}

    for v in dados["veiculos"]:

        caminh[v["placa"]], _ = Caminhao.objects.get_or_create(placa=v["placa"], defaults={k: v[k] for k in ("veiculo", "motorista")})

    criados = 0

    for item in dados["lancamentos"]:

        campos = {k: v for k, v in item.items() if k not in ("placa", "chave_importacao")}

        campos["caminhao"] = caminh[item["placa"]]

        _, novo = LancamentoTransporte.objects.get_or_create(chave_importacao=item["chave_importacao"], defaults=campos)

        criados += int(novo)

    novos = 0

    for item in dados["remuneracoes"]:

        _, novo = RemuneracaoMotorista.objects.get_or_create(caminhao=caminh[item["placa"]], motorista=item["motorista"], competencia=item["competencia"], defaults={k: item[k] for k in ("salario", "comissao", "comissao_percentual", "base_por_aba", "observacoes")})

        novos += int(novo)

    return dict(lancamentos_criados=criados, remuneracoes_criadas=novos, duplicados=len(dados["lancamentos"]) - criados)
